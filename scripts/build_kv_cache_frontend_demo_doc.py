#!/usr/bin/env python3
"""Build the KV Cache frontend demo operator guide as a DOCX."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs/demo/kv-cache/AgenticSRE_KV_Cache前端演示操作流程.docx")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FB"
GREEN = "EAF5EE"
AMBER = "FFF4D6"
RED = "FBEAEA"
INK = "1F2937"
MUTED = "5F6B7A"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1700, 7660])
    shade(table.cell(0, 0), BLUE)
    shade(table.cell(0, 1), fill)
    set_cell_text(table.cell(0, 0), label, bold=True, color="FFFFFF", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), text, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, checked=False):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(("☐ " if checked else "") + text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(INK)


def add_step(doc, number, title, actions, wait, expected, talk):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{number}. {title}")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    meta = doc.add_table(rows=1, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [1700, 7660])
    shade(meta.cell(0, 0), LIGHT_BLUE)
    shade(meta.cell(0, 1), "FFFFFF")
    set_cell_text(meta.cell(0, 0), f"等待 {wait}", bold=True, color=DARK_BLUE, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(meta.cell(0, 1), expected, bold=True, color=INK, size=9.5)
    for action in actions:
        add_bullet(doc, action)
    add_callout(doc, "现场讲解", talk)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.2
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 11.5, DARK_BLUE, 10, 5),
):
    style = styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "AgenticSRE | KV Cache 故障闭环演示"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("现场操作手册 · 2026-07-12")
run.font.name = "Microsoft YaHei"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string(MUTED)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(3)
r = title.add_run("AgenticSRE KV Cache 故障\n前端演示操作流程")
r.bold = True
r.font.name = "Microsoft YaHei"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(25)
r.font.color.rgb = RGBColor.from_string("184A9C")
subtitle = doc.add_paragraph("真实注入 → 多源观测 → 智算事件 → Agent+Skill 诊断 → 受控恢复 → 恢复验证")
subtitle.paragraph_format.space_after = Pt(14)
for r in subtitle.runs:
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MUTED)

add_callout(doc, "演示时长", "建议 6–8 分钟。故障信号等待 30–60 秒，Agent+Skill 诊断通常约 1–2 分钟。", GREEN)
add_callout(doc, "重要约束", "KV Cache Exhaustion 场景自带长上下文压力负载，不要额外勾选“背景负载”；如信号未立即出现，等待并刷新，不要重复注入。", AMBER)

doc.add_heading("演示目标与环境", level=1)
env = doc.add_table(rows=6, cols=2)
env.style = "Table Grid"
set_table_geometry(env, [2700, 6660])
for i, (k, v) in enumerate((
    ("目标集群", "t4-vllm-k8s"),
    ("Namespace", "llm-k8s"),
    ("Deployment / Service", "vllm-server"),
    ("模型", "qwen2.5-coder-3b-k8s"),
    ("故障场景", "KV Cache Exhaustion"),
    ("闭环引擎", "Agent+Skill"),
)):
    shade(env.cell(i, 0), LIGHT_BLUE)
    set_cell_text(env.cell(i, 0), k, bold=True, color=DARK_BLUE, size=10)
    set_cell_text(env.cell(i, 1), v, size=10)

doc.add_heading("九步闭环操作", level=1)
steps = [
    (1, "从拓扑图选择故障目标", ["进入“智算故障实验”，确认集群为 t4-vllm-k8s，点击“刷新依赖图”。", "在依赖图中点击 Model qwen2.5-coder-3b-k8s；也可选择 vllm-server Service 或 Pod。", "在弹出面板的“故障类型”下拉框选择 sw-003 — KV Cache 耗尽。", "点击“使用该目标”，确认 Namespace、Deployment、Service 和场景已自动回填。", "引擎选择 Agent+Skill，取消勾选“背景负载”。"], "10 秒", "目标和 sw-003 / KV Cache Exhaustion 已回填到实验控制台。", "操作人从实时依赖图选择对象，无需手工输入资源名称。拓扑选择只回填目标和场景，真实执行仍经过 Dry-run、预检和确认门禁。"),
    (2, "执行安全预检", ["保持 Dry-run 已勾选，“确认真实执行”未勾选。", "点击“预检”。", "如预检未通过，停止演示并按阻断项检查。"], "10–30 秒", "执行反馈显示“预检通过”。", "默认先生成 Dry-run 命令和恢复预案，目标、权限、工具和恢复路径都通过后才允许真实执行。"),
    (3, "真实注入 KV Cache 压力", ["取消 Dry-run，勾选“确认真实执行”。", "点击“注入”并接受确认。", "记录页面返回的 Fault Run ID。"], "10–30 秒", "执行反馈显示已真实注入，压力负载已启动。", "注入不改变模型权重，也不重启服务，只通过可清理的长上下文请求逐步占用 KV Cache。"),
    (4, "观测故障传播", ["进入“智算观测”，选择 t4-vllm-k8s 和 15M 窗口。", "进入“指标曲线”，保持“业务/vLLM”指标组。", "点击“刷新”；如信号不明显，20–30 秒后再刷新。", "重点展示 KV Cache、等待队列、TTFT P95 和 E2E P95。"], "30–60 秒", "KV Cache 通常 >60%，等待队列 >0，TTFT/E2E 上升；服务仍可能 Ready。", "传播链是 KV Cache 饱和 → 请求排队 → TTFT 和端到端时延上升，属于性能退化而非必然完全中断。"),
    (5, "生成智算事件", ["进入“智算事件”，选择 t4-vllm-k8s、活跃状态和最近 1 小时。", "点击“检测并生成事件”。", "找到“GPU 显存或 KV Cache 压力过高”类事件。", "展开“查看指标、证据和处理状态”。"], "10–30 秒", "活跃事件中包含 KV Cache、队列或时延证据。", "事件由观测窗口生成，Agent 收到的是指标和对象上下文，而不是注入场景 ID 或 Ground Truth。"),
    (6, "启动 Agent+Skill 诊断", ["在目标事件上点击“分析”，进入“分析分支”。", "保留五阶段：上下文检索、假设生成、验证规划、调查取证、推理裁决。", "点击“启动 Agent+Skill 分析”。", "等待自动轮询，必要时点击“刷新过程”。", "完成后可点击“打开 RCA 工作流”展示时间线。"], "60–120 秒", "根因指向 KV Cache 容量压力导致请求排队和时延上升。", "Agent 先生成并排序假设，再选择用于验证高优先级假设的专项 Skill；它不是扫描全部 59 个 Skill。"),
    (7, "展示自愈决策", ["回到“智算事件分析与自愈”，进入“自愈分支”。", "在 Agent+Skill 自愈区域点击“启动 Agent+Skill 自愈”。", "展示风险门禁、Dry-run 恢复建议和验证步骤。"], "20–60 秒", "自愈结果展示受控方案，不是无条件下发高风险命令。", "自愈会继续检查置信度、动作风险和验证条件；本场景的最小影响动作是清理压力负载，不重启健康 vLLM。"),
    (8, "执行受控恢复", ["关闭事件弹窗，返回“智算故障实验”。", "在“当前活跃故障”下拉框选择本次 KV Cache Fault Run，核对故障、目标、已注入时长和 Run ID。", "确认页面已自动回填 sw-003、llm-k8s 和 vllm-server。", "保持 Dry-run 关闭和“确认真实执行”已勾选。", "点击“恢复所选故障”，确认对话框中的 Fault Run ID。"], "10–30 秒", "系统按原始 fault_run_id 执行该次注入保存的恢复契约。", "恢复不再依赖用户当前表单是否选对，而是精确关联原始注入记录，执行最小影响恢复。"),
    (9, "验证恢复", ["返回“智算观测 → 指标曲线 → 业务/vLLM”。", "每 15–30 秒刷新一次，最长等待约 120 秒。", "确认等待队列为 0、KV Cache 回落、TTFT/E2E 不再增长。", "确认 Deployment 副本和 Service Endpoint 仍然可用。"], "30–120 秒", "队列归零、KV Cache 回落、服务保持 Ready，形成闭环。", "闭环完成不是以恢复命令返回成功为准，而是以业务队列、KV Cache、时延和 Ready 状态的重新采样结果为准。"),
]
for args in steps:
    add_step(doc, *args)

doc.add_heading("诊断正确性验收", level=1)
for item in (
    "根因必须明确包含 KV Cache 容量压力/饱和/耗尽导致请求排队和时延上升的语义。",
    "证据至少覆盖 KV Cache 高占用、等待队列增长和 TTFT/E2E 上升中的两类。",
    "关键 Skill 可能为 llm.vllm_scheduler_queue 或 llm.gpu_resource_pressure；以最终根因和证据为准，不硬编码绑定固定 Skill ID。",
    "如结论仅指向 CPU、磁盘或笼统 GPU 故障，且无法解释 KV Cache 和排队证据，则不算诊断成功。",
):
    add_bullet(doc, item)

doc.add_heading("现场异常处理", level=1)
fallbacks = [
    ("注入后指标未上升", "先等待 30 秒并刷新。不要立即重复注入。"),
    ("智算事件未出现", "确认集群和最近 1 小时窗口，再点击“检测并生成事件”。"),
    ("Agent+Skill 仍在运行", "点击“刷新过程”，不要重复启动新 Run。"),
    ("大模型调用超时", "保留当前 Fault Run，先重试分析；仍失败时展示已录制 Demo，然后必须执行“恢复”。"),
    ("恢复后未立即归零", "已接收请求可能仍在排空，持续观测，最长等待约 120 秒。"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
set_table_geometry(table, [3000, 6360])
for j, text in enumerate(("现象", "处理方式")):
    shade(table.cell(0, j), BLUE)
    set_cell_text(table.cell(0, j), text, bold=True, color="FFFFFF", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
for symptom, action in fallbacks:
    cells = table.add_row().cells
    set_table_geometry(table, [3000, 6360])
    shade(cells[0], LIGHT_BLUE)
    set_cell_text(cells[0], symptom, bold=True, color=DARK_BLUE, size=9.2)
    set_cell_text(cells[1], action, size=9.2)

doc.add_heading("最终验收清单", level=1)
for item in (
    "预检通过后才真实注入。",
    "注入后出现 KV Cache、等待队列和 TTFT/E2E 传播信号。",
    "智算事件来自观测窗口，分析输入不包含注入真值。",
    "Agent+Skill 实际经过假设、规划、调查和推理阶段。",
    "最终根因明确指向 KV Cache 容量压力与请求排队。",
    "置信度、关键证据和 Skill 调用过程可见。",
    "自愈阶段展示风险门禁、最小影响动作和验证计划。",
    "恢复后等待队列归零、KV Cache 回落、服务保持 Ready。",
):
    add_bullet(doc, item, checked=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
